import ollama
import os
import subprocess
import shutil
from duckduckgo_search import DDGS
import sys
import io
from pypdf import PdfReader
from docx import Document
import openpyxl
from pptx import Presentation
import pytesseract
from PIL import ImageEnhance, ImageFilter
import json
import re
import cv2

chat_history = [] # 🧠 ذاكرة الـ Agent
current_project_name = ""       # هيحفظ اسم البروجكت الحالي
awaiting_project_name = False   # حالة انتظار اسم البروجكت من اليوزر
pending_task = ""               # لحفظ الطلب الأصلي لحد ما اليوزر يكتب الاسم

pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

sys.stdout.reconfigure(encoding="utf-8", line_buffering=True)
sys.stdin.reconfigure(encoding="utf-8") # 🔥 السطر ده هو اللي هيحل مشكلة الرموز الغريبة في الإدخال!

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
MEMORY_FILE = os.path.join(BASE_DIR, "agent_memory.txt")
CONTEXT_FILE = os.path.join(BASE_DIR, "project_context.json")

# ================= PROJECT CONTEXT =================
project_context = []

last_user_input = ""

active_intent = "default"

# ================= MODE =================

def detect_mode(user):
    t = user.lower()
    if t.startswith("build"):
        return "PROJECT"
    if t == "run":
        return "RUN"
    if t.startswith("delete") or t.startswith("remove"):
        return "DELETE"
    if "job" in t:
        return "JOB"
    if t.startswith("attach"):
        return "ATTACH"
    if t.strip() == "clear":
        return "CLEAR"
    return "CHAT"


# ================= MEMORY =================

def save_last_project(path):
    with open(MEMORY_FILE, "w", encoding="utf-8") as f:
        f.write(path)

def load_last_project():
    if not os.path.exists(MEMORY_FILE):
        return None
    with open(MEMORY_FILE, "r", encoding="utf-8") as f:
        return f.read().strip()

# ================= CONTEXT MEMORY =================

def save_project_context():
    try:
        with open(CONTEXT_FILE, "w", encoding="utf-8") as f:
            json.dump(project_context, f, ensure_ascii=False, indent=2)
    except Exception as e:
        print("⚠ Context save error:", e)


def load_project_context():
    global project_context

    if not os.path.exists(CONTEXT_FILE):
        return

    try:
        with open(CONTEXT_FILE, "r", encoding="utf-8") as f:
            project_context = json.load(f)

        print(f"🧠 Loaded project context ({len(project_context)} files)")

    except Exception as e:
        print("⚠ Context load error:", e)

# ================= HELPERS =================

def is_valid_path(name):
    invalid = ["*", "?", "<", ">", "|", "\""]
    for c in invalid:
        if c in name:
            return False
    return True

def clean_line(line):
    line = line.strip()
    if not line: return None
    if "```" in line: return None
    if ":" in line: return None
    if line.startswith("-") or line.startswith("*"): return None
    if "Here is" in line or "structure" in line: return None # فلترة إضافية
    if not is_valid_path(line): return None
    return line


# ================= BUILD =================

def build_structure(task):
    print("🔨 Designing structure...")
    r = ollama.chat(
        model="llama3",
        messages=[{
            "role":"user",
            "content":f"Return ONLY python project structure for {task}. List files and folders line by line. NO explanation."
        }]
    )

    lines = r["message"]["content"].split("\n")
    files = []

    for raw in lines:
        line = clean_line(raw)
        if not line: continue

        try:
            if line.endswith("/") or line.endswith("\\"):
                os.makedirs(line, exist_ok=True)
                print("📁 Folder:", line)
            else:
                folder = os.path.dirname(line)
                if folder:
                    os.makedirs(folder, exist_ok=True)
                
                # إنشاء ملف فارغ
                with open(line, "w", encoding="utf-8") as f:
                    f.write("")
                
                files.append(line)
                print("📄 File:", line)

        except Exception as e:
            print(f"⚠ Skipped bad path: {line} ({e})")

    # لو مفيش ملفات اتعملت، نعمل ملف رئيسي اجباري
    if len(files) == 0:
        with open("main.py", "w", encoding="utf-8") as f:
            f.write("")
        files.append("main.py")

    return files


# ================= GENERATE =================

def generate_code(files, task):
    for file in files:
        print(f"⚙ Generating code for: {file}...")
        
        prompt = f"Write ONLY python code for {file} in {task}. NO markdown blocks. NO explanation."
        
        r = ollama.chat(
            model="llama3",
            messages=[{"role":"user", "content": prompt}]
        )

        code = r["message"]["content"]
        # تنظيف إضافي للكود
        code = code.replace("```python", "").replace("```", "")
        if code.strip().startswith("Here is"): # لو الذكاء الاصطناعي رغي
            code = "# Generated code\n" + code.split("\n", 1)[1]

        with open(file, "w", encoding="utf-8") as f:
            f.write(code)
        
        print(f"✅ Written: {file}")


# ================= RUN =================

def run_python(file_path):
    print(f"\n🚀 Running: {file_path}")
    folder = os.path.dirname(file_path)
    script = os.path.basename(file_path)

    try:
        subprocess.run(
            [sys.executable, script], # استخدام بايثون الحالي
            timeout=15, # وقت أطول للمدخلات
            cwd=folder
        )
    except subprocess.TimeoutExpired:
        print("⚠ Program timed out (loop or waiting for input).")
    except Exception as e:
        print(f"❌ Run Error: {e}")

def auto_run_and_fix(task_description, initial_code):
    temp_file = "sandbox_test.py"
    current_code = initial_code
    attempts = 0
    max_attempts = 3 # هيحاول يصلح نفسه لحد 3 مرات

    while attempts < max_attempts:
        attempts += 1
        print(f"🧪 Testing code (Attempt {attempts})...")
        
        with open(temp_file, "w", encoding="utf-8") as f:
            f.write(current_code)

        try:
            result = subprocess.run(
                [sys.executable, temp_file],
                capture_output=True,
                text=True,
                timeout=5 # وقت قليل عشان لو فيه Infinite Loop
            )

            if result.returncode == 0:
                print(f"✅ Success on attempt {attempts}!")
                return current_code, result.stdout
            
            # لو فشل، نبعت الـ Error للموديل
            print(f"❌ Attempt {attempts} failed. Refactoring...")
            error_msg = result.stderr
            
            fix_prompt = f"""FIX THIS PYTHON CODE.
            TASK: {task_description}
            ERROR: {error_msg}
            CODE TO FIX:
            {current_code}
            
            RULES: Return ONLY the raw code. No markdown, no explanations."""
            
            response = ollama.chat(model="llama3", messages=[{"role": "user", "content": fix_prompt}])
            raw_content = response["message"]["content"]
            
            # تنظيف صارم للكود من أي رغي جانبي
            current_code = raw_content.replace("```python", "").replace("```", "").strip()
            if "import" not in current_code and "print" not in current_code: # حماية لو رد بكلام رغي
                 current_code = initial_code # ارجع للأصل لو الموديل خرف

        except subprocess.TimeoutExpired:
            return current_code, "❌ Execution timed out (Possible infinite loop)."
        except Exception as e:
            return current_code, f"❌ System Error: {str(e)}"

    return current_code, "⚠️ Could not fix code after 3 attempts."
    
# ================= DELETE =================

def delete_tool(task):
    name = task.replace("delete", "").replace("remove", "").strip()
    
    # محاولة الحذف من المسار الحالي
    target = os.path.abspath(name)
    
    if os.path.exists(target):
        try:
            shutil.rmtree(target)
            print(f"🗑️ Deleted: {target}")
            
            # تنظيف الذاكرة
            last = load_last_project()
            if last and name in last:
                if os.path.exists(MEMORY_FILE):
                    os.remove(MEMORY_FILE)
        except Exception as e:
            print(f"❌ Error deleting: {e}")
    else:
        print(f"❌ Folder '{name}' not found.")


# ================= PROJECT =================

def project_tool(task):
    folder_name = task.replace("build", "").strip().replace(" ", "_").lower()
    if not folder_name: folder_name = "project"

    # إنشاء المشروع في المكان الحالي اللي اليوزر واقف فيه
    project_path = os.path.join(os.getcwd(), folder_name)
    
    if os.path.exists(project_path):
        print(f"⚠ Folder '{folder_name}' already exists. Updating inside it.")
    
    os.makedirs(project_path, exist_ok=True)
    print(f"📁 Project Path: {project_path}")

    cwd = os.getcwd()
    os.chdir(project_path)

    files = build_structure(task)
    generate_code(files, task)

    # اختيار ملف التشغيل بذكاء
    chosen = None
    priority = ["main.py", "app.py", "run.py", f"{folder_name}.py"]
    
    for p in priority:
        if p in files:
            chosen = p
            break
            
    if not chosen and files:
        chosen = files[0]

    if chosen:
        abs_path = os.path.abspath(chosen)
        save_last_project(abs_path)
        
        # العودة للمسار الأصلي قبل التشغيل لتجنب الأخطاء
        os.chdir(cwd)
        run_python(abs_path)
        print("🔥 Project completed.")
    else:
        print("❌ No files generated to run.")
        os.chdir(cwd)

# ================= JOB =================
import warnings
import urllib.parse
import webbrowser
warnings.filterwarnings("ignore", category=RuntimeWarning)

def job_tool(task):
    print("🧠 Analyzing job market...")
    
    # 🔥 Prompt صارم جداً مع مثال واضح
    prompt = f"""Extract job details from this request: '{task}'. 
    Default Title: Game Developer OR Python AI
    Default Platform: Upwork
    Default Location: Remote OR Egypt
    Reply ONLY with this exact format: TITLE | PLATFORM | LOCATION
    Example: Game Developer | Upwork | Remote
    CRITICAL RULE: DO NOT write 'Here are the details', do not use markdown like **, and DO NOT add any conversational text. JUST the raw format."""
    
    try:
        r = ollama.chat(model="llama3", messages=[{"role": "user", "content": prompt}])
        ans = r["message"]["content"].strip()
        
        # 🛡️ فلتر أمان بايثون (عشان لو الموديل رغى برضه، نقص كلامه)
        if ":" in ans:
            ans = ans.split(":")[-1] # لو كتب نقطتين، بناخد اللي بعدهم بس
        ans = ans.replace('**', '').replace('"', '').strip()
        
        parts = ans.split('|')
        title = parts[0].strip() if len(parts) > 0 else "Game Developer"
        platform = parts[1].strip().lower() if len(parts) > 1 else "upwork"
        location = parts[2].strip() if len(parts) > 2 else "Remote"
        
        # تنظيف أخير للعنوان لو لسه فيه أي شوائب
        if title.lower().startswith("here are") or title.lower().startswith("extracted"):
            title = "Game Developer OR Python AI"
            
        print(f"🚀 Opening {platform.title()} for {title} roles...")
        
        title_encoded = urllib.parse.quote(title)
        
        if "upwork" in platform:
            url = f"https://www.upwork.com/nx/search/jobs/?q={title_encoded}&sort=recency"
        else:
            loc_encoded = urllib.parse.quote(location)
            url = f"https://www.linkedin.com/jobs/search/?keywords={title_encoded}&location={loc_encoded}"
            
        webbrowser.open(url)
        print("✅ Browser opened with real-time job listings!")
        
    except Exception as e:
        print(f"❌ Error setting up job search: {e}")

# ================= CHAT (FULLY AUTONOMOUS) =================

def self_correct(draft_response, user_query):
    #print("🧠 Reflection Layer: Reviewing and polishing the response...")
    
    # رسالة تظهر وقت الـ (Idle) عشان تفهم إن التأخير ده بسبب تحميل الموديل
    #print("⏳ Swapping models in memory (Loading Llama 3)...", flush=True) 
    
    prompt = f"""You are a professional Content Editor. You are reviewing a draft response from an AI assistant.
    
    USER QUERY: "{user_query}"
    DRAFT RESPONSE: "{draft_response}"

    YOUR STRICT MISSION:
    1. Clean up the language and make it sound natural and professional.
    2. Keep the core information and meaning of the draft EXACTLY as it is. Do NOT fabricate or add new information.
    3. If the draft describes files, images, or code, DO NOT say you cannot see them. Trust the draft completely.
    4. NEVER apologize or say "I am an AI", "As an AI", or "I have limitations". 
    5. Just return the polished, final version of the response directly.

    FINAL OUTPUT ONLY. NO INTRODUCTIONS. NO CHATTY TEXT."""
    
    try:
        # 🔥 تفعيل خاصية الـ Streaming
        stream = ollama.chat(model="llama3", messages=[{"role": "user", "content": prompt}], stream=True)
        
        # بنطبع البداية ونجبر الشاشة تطلعها فوراً
        sys.stdout.write("\n🤖 Agent: ")
        sys.stdout.flush()
        
        final_reply = ""
        for chunk in stream:
            word = chunk['message']['content']
            # بنكتب الكلمة ونجبر الشاشة تعرضها في نفس اللحظة (غصب عن التيرمنال)
            sys.stdout.write(word)
            sys.stdout.flush()
            final_reply += word
            
        sys.stdout.write("\n\n")
        sys.stdout.flush()
        
        return final_reply
        
    except Exception as e:
        print(f"\n⚠️ Reflection failed: {e}")
        return draft_response

def detect_programming_domain(code):
    code_lower = code.lower()

    if "uclass" in code_lower or "generated_body" in code_lower:
        return "unreal"

    if "using unityengine" in code_lower or "monobehaviour" in code_lower:
        return "unity"

    if "import " in code_lower or "def " in code_lower:
        return "python"

    if "#include" in code_lower:
        return "cpp"

    if "using system" in code_lower:
        return "csharp"

    return "unknown"

def validate_unreal(code):
    required = ["#include", "GENERATED_BODY"]
    for r in required:
        if r.lower() not in code.lower():
            return False, f"Missing Unreal requirement: {r}"

    if ".generated.h" not in code:
        return False, "Missing .generated.h include"

    # regex صح عشان مينفعش يعتبر UActorComponent كـ UComponent أو UObject
    if re.search(r'public\s+UComponent\b', code):
        return False, "Invalid Unreal inheritance: Use UActorComponent instead of UComponent"

    if re.search(r'public\s+UObject\b', code):
        return False, "Component must inherit from UActorComponent, not UObject"

    return True, "Valid Unreal Header"

def validate_unity(code):
    if "MonoBehaviour" not in code:
        return False, "Unity script must inherit from MonoBehaviour"

    if "public class" not in code:
        return False, "Unity script missing public class"

    if "void Start" not in code and "void Update" not in code:
        return False, "Unity script missing Start or Update method"

    return True, "Valid Unity Code"

def validate_python(code, task):
    fixed_code, output = auto_run_and_fix(task, code)

    if "Could not fix" in output:
        return False, output

    return True, fixed_code

def validate_code(code, task):
    domain = detect_programming_domain(code)

    if domain == "unreal":
        return validate_unreal(code)

    if domain == "unity":
        return validate_unity(code)

    if domain == "python":
        return validate_python(code, task)

    return True, "Generic code assumed valid"

def detect_requested_language(task):
    t = task.lower()

    if "unreal" in t:
        return "unreal"

    if "unity" in t or "يونتي" in t:
        return "unity"

    if "python" in t or "بايثون" in t:
        return "python"

    return "generic"

LANGUAGE_RULES = {
    "unreal": {
        "type": "multi",
        "extensions": [".h", ".cpp"],
    },
    "unity": {
        "type": "single",
        "extensions": [".cs"],
    },
    "python": {
        "type": "single",
        "extensions": [".py"],
    },
    "cpp": {
        "type": "single",
        "extensions": [".cpp"],
    },
    "csharp": {
        "type": "single",
        "extensions": [".cs"],
    },
}

def auto_fix_unreal_header(header_code):
    """
    Inject missing .generated.h include if absent.
    """
    if ".generated.h" in header_code:
        return header_code, False

    # Extract class name
    class_match = re.search(
        r"class\s+(?:[A-Za-z0-9_]+\s+)?([A-Za-z_][A-Za-z0-9_]*)",
        header_code
    )

    if not class_match:
        return header_code, False

    class_name = class_match.group(1)

    include_line = f'#include "{class_name}.generated.h"\n'

    # نحاول نحطه بعد آخر include
    includes = list(re.finditer(r'#include\s+["<].*[">]', header_code))
    if includes:
        last_include = includes[-1]
        insert_pos = last_include.end()
        fixed_header = (
            header_code[:insert_pos]
            + "\n"
            + include_line
            + header_code[insert_pos:]
        )
    else:
        # لو مفيش includes خالص
        fixed_header = include_line + header_code

    return fixed_header, True
    
def extract_and_save_scripts(text, project_name):

    pattern = r"```(?:[a-zA-Z0-9+#]*)\n?(.*?)```"
    matches = re.findall(pattern, text, re.DOTALL | re.IGNORECASE)

    if not matches:
        return False

    safe_project_name = re.sub(r'[\\/*?:"<>|]', "", project_name).strip().replace(" ", "_")
    folder_name = os.path.join("Generated_Scripts", safe_project_name)
    os.makedirs(folder_name, exist_ok=True)

    # ==========================================
    # 🔥 Detect Domain From First Block
    # ==========================================
    first_domain = detect_programming_domain(matches[0])

    if first_domain in LANGUAGE_RULES:
        rule = LANGUAGE_RULES[first_domain]

        # ==========================================
        # 🚀 MULTI-FILE (Unreal)
        # ==========================================
        if rule["type"] == "multi":

            if len(matches) < 2:
                print("❌ Unreal requires header and cpp blocks.")
                return False

            header_code = matches[0].strip()
            cpp_code = matches[1].strip()

            # 🔒 Validate Header
            is_valid_header, msg_header = validate_unreal(header_code)

            if not is_valid_header and "generated.h" in msg_header:

                print("⚠ Missing .generated.h — attempting auto-fix...")

                header_code, fixed = auto_fix_unreal_header(header_code)

                if fixed:
                    # Re-validate after fix
                    is_valid_header, msg_header = validate_unreal(header_code)

            if not is_valid_header:
                print(f"❌ Validation Failed: {msg_header}")
                return False

            # 🔒 Validate CPP
            if "::" not in cpp_code:
                print("❌ Unreal CPP missing implementation (:: not found)")
                return False

            # ==========================================
            # 🔥 Structural Validation
            # ==========================================

            # كل أسماء الكلاسات الموجودة في الـ CPP (عشان نتجاهل Constructors بأي اسم)
            cpp_class_names = set(re.findall(r'([A-Za-z_][A-Za-z0-9_]*)::', cpp_code))

            # Overrides معروفة في Unreal
            known_overrides = {
                "BeginPlay", "TickComponent", "EndPlay",
                "InitializeComponent", "GetLifetimeReplicatedProps",
                "SetupInputComponent", "PostInitializeComponents"
            }

            # تنظيف الـ Macros قبل استخراج الـ declarations
            header_clean = re.sub(r'\b(UFUNCTION|UPROPERTY|UCLASS|USTRUCT|UENUM)\s*\([^)]*\)', '', header_code)
            declared_funcs = set(re.findall(
                r'\b([A-Za-z_][A-Za-z0-9_]*)\s*\([^)]*\)\s*(?:const\s*)?;',
                header_clean
            ))

            # استخراج الـ implementations من الـ CPP
            implemented_funcs = re.findall(r'::\s*([A-Za-z_][A-Za-z0-9_]*)\s*\(', cpp_code)

            for func in implemented_funcs:
                # تجاهل Constructors/Destructors (اسمهم = اسم الكلاس)
                if func in cpp_class_names or func.lstrip('~') in cpp_class_names:
                    continue
                # تجاهل Overrides المعروفة
                if func in known_overrides:
                    continue
                if func not in declared_funcs:
                    print(f"⚠ Note: '{func}' not in header — treating as override, continuing.")
            
            # 🔍 Extract Class Name From Header
            class_match = re.search(
                r"class\s+(?:[A-Za-z0-9_]+\s+)?([A-Za-z_][A-Za-z0-9_]*)",
                header_code
            )
            class_name = class_match.group(1) if class_match else "UnrealClass"

            header_path = os.path.join(folder_name, f"{class_name}.h")
            cpp_path = os.path.join(folder_name, f"{class_name}.cpp")

            with open(header_path, "w", encoding="utf-8") as f:
                f.write(header_code)

            with open(cpp_path, "w", encoding="utf-8") as f:
                f.write(cpp_code)

            print(f"\n💾 Script saved -> {header_path}")
            print(f"💾 Script saved -> {cpp_path}")

            return True

    # ==========================================
    # 📁 SINGLE FILE LANGUAGES
    # ==========================================
    all_valid = True

    for i, code in enumerate(matches):
        code = code.strip()

        is_valid, message = validate_code(code, project_name)
        if not is_valid:
            print(f"\n❌ Validation Failed: {message}")
            all_valid = False
            continue

        domain = detect_programming_domain(code)

        if domain not in LANGUAGE_RULES:
            extension = ".txt"
        else:
            rule = LANGUAGE_RULES[domain]
            extension = rule["extensions"][0]

        # 🔍 Extract Class Name
        class_match = re.search(
            r"class\s+(?:[A-Za-z0-9_]+\s+)?([A-Za-z_][A-Za-z0-9_]*)",
            code
        )
        class_name = class_match.group(1) if class_match else f"Script_{i}"

        file_name = f"{class_name}{extension}"
        file_path = os.path.join(folder_name, file_name)

        with open(file_path, "w", encoding="utf-8") as f:
            f.write(code)

        print(f"\n💾 Script saved -> {file_path}")

    return all_valid

def chat_tool(task):
    global active_intent, chat_history, current_project_name, awaiting_project_name, pending_task
    
    if task:
        task = task.strip()
        
    if task == "Analyze and describe the attached files in detail.":
        task = ""

    # =========================================================
    if awaiting_project_name and task:
        current_project_name = task
        awaiting_project_name = False
        print(f"\n🤖 Agent: عظيم! تم تحديد اسم البروجكت: '{current_project_name}'. جاري كتابة السكريبت...\n")
        task = pending_task 
    # =========================================================

    if not task:
        count = len(project_context)
        file_word = "file" if count == 1 else "files"
        print(f"📎 Successfully attached {count} {file_word} to project context!")
        if count >= 2:
            print("Tell me what you want to do with them (e.g., compare, summarize).")
        else:
            print("Ask a specific question about this file.")
        active_intent = "default"
        return

    intent = detect_intent(task)
    active_intent = intent
    working_context = select_relevant_files(task, project_context)
    
    images = [item["path"] for item in working_context if item["type"] == "image"]
    videos = [item["path"] for item in working_context if item["type"] == "video"]
    text_context = ""
    
    for item in working_context:
        if item["type"] not in ["image", "video"]:
            text_context += f"\n--- File: {os.path.basename(item['path'])} ---\n"
            if "data" in item:
                text_context += json.dumps(item["data"])
            else:
                try:
                    with open(item["path"], "r", encoding="utf-8", errors="ignore") as f:
                        text_context += f.read(1500)
                except:
                    pass

    model_name = "llava" if images else "llama3"
    
    status_parts = []
    if images: status_parts.append(f"{len(images)} images")
    if videos: status_parts.append(f"{len(videos)} videos")
    
    other_files_count = len(working_context) - len(images) - len(videos)
    if other_files_count > 0:
        status_parts.append(f"{other_files_count} text/data files")

    if status_parts:
        report = " and ".join(status_parts)
        print(f"🤔 Thinking about {report}...")

    # 3. توجيه الذكاء الاصطناعي بناءً على النية اللي اكتشفها
    if intent == "compare":
        if videos:
            system_prompt = "You are an expert video analyst. Compare the attached videos explicitly."
        elif images:
            system_prompt = "You are an expert visual analyst. Compare the attached images explicitly."
        else:
            system_prompt = "You are an expert data analyst. Compare the attached files explicitly."
    elif intent == "summarize":
        system_prompt = "You are an expert summarizer. Provide a concise, highly accurate summary of the provided files."
    elif intent == "detect_issues":
        system_prompt = "You are an expert debugger and reviewer. Analyze the provided files to find any bugs, errors, or issues."
    else:
        if len(working_context) > 0:
            system_prompt = "You are an elite Game Dev & AI assistant. Answer accurately based on the attached files."
        else:

            language = detect_requested_language(task)

            if language == "unity":
                system_prompt = """
                You are a senior Unity C# developer.
                Generate ONLY valid Unity C# scripts.
                Must inherit from MonoBehaviour.
                No Unreal code.
                No explanations outside code blocks.
                """

            elif language == "unreal":
                system_prompt = """You are a senior Unreal Engine C++ developer.

STRICT OUTPUT RULES:
1) Return EXACTLY TWO ```cpp blocks. Nothing else outside them.
2) First block = Header (.h), Second block = CPP (.cpp)
3) Header MUST:
   - Start with #pragma once
   - #include "CoreMinimal.h"
   - #include "Components/ActorComponent.h"
   - #include "<ExactClassName>.generated.h"
   - UCLASS(ClassGroup=(Custom), meta=(BlueprintSpawnableComponent))
   - class MYGAME_API <ClassName> : public UActorComponent
   - GENERATED_BODY()
   - Declare constructor, ALL functions, UPROPERTY, UFUNCTION, delegates
4) CPP MUST:
   - #include "<ExactClassName>.h" (SAME name as the class file)
   - Implement constructor and ALL declared functions
5) NO text or explanations outside the two code blocks.

EXAMPLE:
```cpp
#pragma once
#include "CoreMinimal.h"
#include "Components/ActorComponent.h"
#include "MyComp.generated.h"

UCLASS(ClassGroup=(Custom), meta=(BlueprintSpawnableComponent))
class MYGAME_API UMyComp : public UActorComponent
{
    GENERATED_BODY()
public:
    UMyComp();
    UPROPERTY(EditAnywhere, BlueprintReadWrite, Category="Stats")
    float Value;
    UFUNCTION(BlueprintCallable)
    void DoSomething(float Amount);
protected:
    virtual void BeginPlay() override;
};
```
```cpp
#include "MyComp.h"

UMyComp::UMyComp()
{
    PrimaryComponentTick.bCanEverTick = false;
}
void UMyComp::BeginPlay()
{
    Super::BeginPlay();
}
void UMyComp::DoSomething(float Amount)
{
    Value -= Amount;
}
```"""

            elif language == "python":
                system_prompt = """
                You are a senior Python developer.
                Generate clean runnable Python code only.
                No markdown explanations.
                """

            else:
                system_prompt = """
                You are a senior software engineer.
                Generate correct programming code only.
                No explanations outside code blocks.
                """

    if text_context:
        task = f"Context from text files:\n{text_context}\n\nUser Task: {task}"

    try:
        chat_history.append({"role": "user", "content": task})
        # نحافظ على آخر 10 رسايل بس عشان منثقلش الموديل
        if len(chat_history) > 10:
            chat_history = chat_history[-10:]
        messages_to_send = [{"role": "system", "content": system_prompt}] + chat_history

        if not text_context and not images and not videos:
            is_unity = any(word in task.lower() for word in ["unity", "c#", "combat", "game"])
            is_python = any(word in task.lower() for word in ["python", "بايثون", "script"])

            if is_python and not is_unity and any(word in task.lower() for word in ["code", "برنامج", "كود"]):
                r = ollama.chat(model=model_name, messages=messages_to_send)
                initial_code = r["message"]["content"].replace("```python", "").replace("```", "").strip()
                
                fixed_code, output = auto_run_and_fix(task, initial_code)
                
                chat_history.append({"role": "assistant", "content": f"```python\n{fixed_code}\n```"})
                print(f"\n🤖 Agent (Python Verified):\n```python\n{fixed_code}\n```\n📝 Output: {output}\n")
            
            else:
                is_script_request = any(word in task.lower() for word in [
                    "script", "code", "سكريبت", "كود", "system", "برمج",
                    "اكتب", "write", "component", "class"
                ])
                
                if is_script_request:
                    if not current_project_name:
                        print("🤔 Checking project name...")
                        name_check_prompt = f"Extract the game or project name from this text. If none is mentioned, reply ONLY with 'NONE'. Text: '{task}'"
                        r = ollama.chat(model="llama3", messages=[{"role": "user", "content": name_check_prompt}])
                        extracted_name = r["message"]["content"].strip()
                        
                        if "NONE" not in extracted_name.upper() and len(extracted_name) < 30:
                            current_project_name = extracted_name
                        else:
                            awaiting_project_name = True
                            pending_task = task
                            chat_history.pop()
                            print("\n🤖 Agent: حلو جداً! بس قبل ما أكتب الكود، إيه اسم اللعبة أو البروجكت بتاعك عشان أعمله فولدر مخصوص؟\n")
                            sys.stdout.flush()
                            return
                    
                    r = ollama.chat(model=model_name, messages=messages_to_send)
                    full_response = r["message"]["content"]

                    success = extract_and_save_scripts(full_response, current_project_name)

                    if success:
                        clean_text = f"""عاش يا هندسة! 🫡 الأكواد اتبرمجت واتقسمت صح.

                    [ 💾 تم حفظ الملفات بنجاح في: Generated_Scripts/{current_project_name.replace(' ', '_')} ]
                    """
                        current_project_name = ""  # reset عشان المشروع الجاي ياخد اسمه الصح
                    else:
                        clean_text = "⚠️ تم إيقاف الحفظ بسبب أخطاء في الكود."

                    print(f"\n🤖 Agent: {clean_text}\n\n")
                    print("🏁 Done.", flush=True)
                    chat_history.append({"role": "assistant", "content": clean_text})
                    
                else:
                    stream = ollama.chat(model=model_name, messages=messages_to_send, stream=True)
                    sys.stdout.write("\n🤖 Agent: ")
                    sys.stdout.flush()
                    
                    full_response = ""
                    for chunk in stream:
                        word = chunk['message']['content']
                        full_response += word
                        sys.stdout.write(word)
                        sys.stdout.flush()
                    sys.stdout.write("\n\n")
                    sys.stdout.flush()
                    print("🏁 Done.", flush=True)
                    chat_history.append({"role": "assistant", "content": full_response})
        
        else:
            chat_history[-1]["images"] = images if images else None
            messages_with_files = [{"role": "system", "content": system_prompt}] + chat_history
            r = ollama.chat(model=model_name, messages=messages_with_files)
            result = r["message"]["content"]
            chat_history.append({"role": "assistant", "content": result})
            print(f"\n🤖 Agent: {result}\n")
            print("🏁 Done.", flush=True)
            
        sys.stdout.flush()

    except Exception as e:
        print(f"❌ Error: {e}")
        if chat_history:
            chat_history.pop()

# ================= FILE TYPE DETECTOR =================

def detect_file_type(file_path):

    ext = os.path.splitext(file_path)[1].lower()

    code_ext = [
        ".py",".js",".cs",".cpp",".c",".java",
        ".html",".css",".json",".xml",".ts"
    ]

    image_ext = [".png",".jpg",".ico",".jpeg",".webp",".bmp"]
    video_ext = [".mp4",".avi",".mov",".mkv",".webm"]
    pdf_ext = [".pdf"]
    word_ext = [".docx"]
    excel_ext = [".xlsx"]
    ppt_ext = [".pptx"]

    if ext in code_ext:
        return "code"

    if ext in image_ext:
        return "image"

    if ext in video_ext:
        return "video"
    
    if ext in pdf_ext:
        return "pdf"
    
    if ext in word_ext:
        return "word"

    if ext in excel_ext:
        return "excel"

    if ext in ppt_ext:
        return "ppt"

    return "text"

# ================= SMART HANDLERS =================

def detect_intent(user_text):
    text_lower = user_text.lower()
    
    # 🔥 1. كبرنا شبكة الدردشة عشان تشمل الهزار والأسئلة العامة
    default_words = [
        "help", "make", "create", "build", "unity", "code", "write", "game", 
        "how", "what", "can you", "project", "budget", "develop",
        "joke", "tell", "say", "yes", "no", "thanks", "who", "why", "where", "when",
        "ازيك", "عامل", "هلو", "اهلا", "مرحبا", "hello", "hi", "hey", "good"
    ]
    if any(w in text_lower for w in default_words):
        return "default"
        
    compare_words = ["فرق", "قارن", "مقارنة", "اختلاف", "compare", "difference"]
    if any(w in text_lower for w in compare_words):
        return "compare"
        
    summarize_words = ["لخص", "الخلاصة", "باختصار", "الملخص", "summarize", "summary"]
    if any(w in text_lower for w in summarize_words):
        return "summarize"
        
    issue_words = ["مشكلة", "غلطة", "خطأ", "ايرور", "bug", "issue", "error", "fix"]
    if any(w in text_lower for w in issue_words):
        return "detect_issues"
        
    describe_words = ["اشرح", "ايه ده", "تفاصيل", "وصف", "فيها ايه", "describe", "explain"]
    if any(w in text_lower for w in describe_words):
        return "describe"
        
    # 🔥 2. قفلنا ثغرة التفكير الغبي للموديل
    try:
        prompt = f"Categorize into ONE: 'compare', 'summarize', 'detect_issues', 'describe', 'default'. RULE: For coding, general questions, or jokes, YOU MUST pick 'default'. User Request: '{user_text}'. Reply with one word."
        r = ollama.chat(model="llama3", messages=[{"role": "user", "content": prompt}])
        ans = r["message"]["content"].strip().lower()
        
        # لو الموديل جاب سيرة default في كلامه، نعتبرها دردشة فوراً ومندورش على الباقي
        if "default" in ans:
            return "default"
            
        for valid in ["compare", "summarize", "detect_issues", "describe"]:
            if valid in ans:
                return valid
    except:
        pass
        
    return "default"

def select_relevant_files(user_text, context_list):
    total_files = len(context_list)
    if not user_text or total_files <= 1:
        return context_list
        
    print("🧠 Context Layer: Selecting the best files for this task...")
    text_lower = user_text.lower()
    
    # 🔥 فلترة بايثون السريعة: لو طلب "كل" الصور أو كتب رقم بيساوي عدد الملفات
    if "كل" in text_lower or "all" in text_lower or str(total_files) in text_lower:
        print(f"🎯 Selected ALL {total_files} relevant files (Fast Rule)")
        return context_list
        
    files_info = ""
    for i, item in enumerate(context_list, 1): 
        files_info += f"File {i}: {os.path.basename(item['path'])}\n"
        
    prompt = f"""Extract the requested file numbers from the user's text. 
    The user might use Arabic or English ordinal numbers (e.g., الأولى = 1, الرابعة = 4).
    Available Files (Total: {total_files}):
    {files_info}
    User Request: "{user_text}"
    Reply ONLY with the digits separated by commas (e.g., 1, 4). No text."""
    
    try:
        r = ollama.chat(model="llama3", messages=[{"role": "user", "content": prompt}])
        ans = r["message"]["content"].strip()
        
        selected_ids = [int(s) - 1 for s in ans.replace(',', ' ').split() if s.isdigit()]
        selected_files = [context_list[i] for i in selected_ids if 0 <= i < total_files]
        
        if selected_files:
            print(f"🎯 Selected {len(selected_files)} relevant files out of {total_files}")
            return selected_files
            
    except Exception as e:
        print(f"⚠️ Context selection error: {e}")
        
    return context_list

def brain_prompt(file_type, file_name, intent="default"):

    file_name = file_name.lower()

    # 🎯 Intent-based override
    if intent == "compare":
        return "Compare this file with other related files and highlight similarities and differences."

    if intent == "detect_issues":
        return "Analyze this file carefully and detect any problems, inconsistencies, or issues."

    if intent == "summarize":
        return "Provide a clear and concise summary of this content."

    if intent == "describe":
        return "Describe this content in detail with important observations."
    
    # صور
    if file_type == "image":
        return "Describe this image in detail. Mention important visual details."

    # كود
    if file_type == "code":
        return "Explain this code briefly. Detect bugs and suggest improvements."

    # PDF
    if file_type == "pdf":
        if "invoice" in file_name or "receipt" in file_name:
            return "Extract invoice information in structured form."
        return "Summarize this PDF clearly."

    # Word / Text
    if file_type in ["word", "text"]:
        return "Summarize this document and extract key points."

    # Excel
    if file_type == "excel":
        return "Analyze this spreadsheet and explain the data."

    # PowerPoint
    if file_type == "ppt":
        return "Summarize this presentation slide by slide."

    return "Analyze this file intelligently."

def handle_code_file(file_path):

    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read(4000)

        print("🧠 Analyzing code with AI...")

        r = ollama.chat(
            model="llama3",
            messages=[{
                "role":"user",
                "content":f"Explain this code briefly and detect problems:\n\n{content}"
            }]
        )

        print("\n🤖 AI Analysis:")
        print(r["message"]["content"])

    except Exception as e:
        print(f"❌ Code analysis error: {e}")


def handle_image_file(file_path):
    try:
        size = os.path.getsize(file_path) / (1024*1024)
        print(f"🖼 Image attached: {os.path.basename(file_path)}")
        print(f"📦 Size: {size:.2f} MB")
        
        # 🔥 التعديل هنا: رسالة توضيحية فقط
        print("⚡ Vision ready: Ask me to describe it!") 

    except Exception as e:
        print(f"❌ Image error: {e}")


def handle_video_file(file_path):
    print("🎬 Processing video for RTX 3060 Ti...")
    
    try:
        vidcap = cv2.VideoCapture(file_path)
        
        # معلومات الفيديو
        total_frames = int(vidcap.get(cv2.CAP_PROP_FRAME_COUNT))
        fps = vidcap.get(cv2.CAP_PROP_FPS)
        duration = total_frames / fps
        
        print(f"⏱ Duration: {duration:.1f}s | Total Frames: {total_frames}")

        # ✅ استراتيجية 3060 Ti:
        # هناخد 5 لقطات موزعة بالتساوي على الفيديو كله
        # ده أفضل من أخذ لقطة كل ثانية عشان الذاكرة متتخنقش
        target_frame_count = 5 
        
        # لو الفيديو قصير جدا (أقل من 5 ثواني) هناخد لقطة كل ثانية
        if duration < 5:
            step = int(fps) # كل ثانية
        else:
            step = int(total_frames / target_frame_count)

        if step == 0: step = 1

        count = 0
        extracted_count = 0
        
        while True:
            # نقفز للفريم المطلوب مباشرة لتوفير المعالجة
            # (بدل ما نقرأ فريم فريم)
            frame_id = extracted_count * step
            if frame_id >= total_frames:
                break
                
            vidcap.set(cv2.CAP_PROP_POS_FRAMES, frame_id)
            success, image = vidcap.read()
            
            if not success: break
            
            # 🔥 أهم خطوة للـ 8GB VRAM: تصغير الصورة
            # بنخلي الارتفاع 512 بيكسل بس، والعرض يتضبط أوتوماتيك
            # ده بيخلي الصورة خفيفة جداً على الموديل
            height, width = image.shape[:2]
            max_height = 512
            if height > max_height:
                scale = max_height / height
                new_width = int(width * scale)
                image = cv2.resize(image, (new_width, max_height))

            # الحفظ والإضافة
            frame_path = f"{file_path}_frame_{extracted_count}.jpg"
            cv2.imwrite(frame_path, image)
            
            if not any(item["path"] == frame_path for item in project_context):
                project_context.append({
                    "path": frame_path,
                    "type": "image"
                })

            
            print(f"📸 Frame {extracted_count+1}: {os.path.basename(frame_path)} (Resized)")
            extracted_count += 1
            
            # أمان أخير
            if extracted_count >= 6:
                break
        
        save_project_context()
        print(f"⚡ Done! {extracted_count} frames ready. Ask: 'Describe video'")
        
    except Exception as e:
        print(f"❌ Video error: {e}")

def handle_pdf_file(file_path):

    try:
        print("📄 Reading PDF...")

        # نحاول نقرأه كنص عادي الأول
        reader = PdfReader(file_path)

        text = ""

        for page in reader.pages[:3]:
            text += page.extract_text() or ""

        # لو مفيش نص → نستخدم OCR
        if not text.strip():

            print("🧠 No text found, using OCR...")

            from pdf2image import convert_from_path

            images = convert_from_path(
                file_path,
                first_page=1,
                last_page=2,
                poppler_path=r"C:\poppler\Library\bin"
            )

            for img in images:

                # تحويل رمادي
                img = img.convert("L")

                # زيادة التباين
                enhancer = ImageEnhance.Contrast(img)
                img = enhancer.enhance(2)

                # شوية sharpen
                img = img.filter(ImageFilter.SHARPEN)

                text += pytesseract.image_to_string(
                    img,
                    lang="ara+eng",
                    config="--psm 6"
                )

        if not text.strip():
            print("❌ OCR failed.")
            return

        print("🧠 Sending PDF content to AI...")

        r = ollama.chat(
            model="llama3",
            messages=[{
                "role":"user",
                "content": f"""
                Extract ONLY real information from this document.

                Return STRICT JSON format only:

                {{
                "company": "",
                "date": "",
                "total_amount": "",
                "invoice_number": "",
                "details": ""
                }}

                Do NOT guess.
                If something missing leave it empty.

                Document:
                {text[:4000]}"""
            }]
        )

        print("\n🤖 PDF Summary:")
        print(r["message"]["content"])

        try:
            raw = r["message"]["content"]

            # استخراج JSON من أي كلام حوالينه
            match = re.search(r"\{.*\}", raw, re.DOTALL)

            if match:
                json_text = match.group(0)
                data = json.loads(json_text)

                print("✅ JSON Parsed Successfully")
                print(data)

                # حفظ البيانات داخل project context
                for item in project_context:
                    if item["path"] == file_path:
                        item["data"] = data
                        break

                save_project_context()

            else:
                print("⚠ No JSON found.")

        except Exception as e:
            print("⚠ JSON parse error:", e)

    except Exception as e:
        print(f"❌ PDF error: {e}")

def handle_word_file(file_path):

    try:
        print("📄 Reading Word file...")

        doc = Document(file_path)

        text = ""
        for p in doc.paragraphs:
            text += p.text + "\n"

        text = text[:4000]

        print("🧠 Sending Word content to AI...")

        r = ollama.chat(
            model="llama3",
            messages=[{
                "role":"user",
                "content":f"Summarize this document:\n\n{text}"
            }]
        )

        print("\n🤖 Word Summary:")
        print(r["message"]["content"])

    except Exception as e:
        print(f"❌ Word error: {e}")


def handle_excel_file(file_path):

    try:
        print("📊 Reading Excel file...")

        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet = wb.active

        rows = []

        for row in sheet.iter_rows(values_only=True):
            rows.append(str(row))

        text = "\n".join(rows)[:4000]

        print("🧠 Sending Excel content to AI...")

        r = ollama.chat(
            model="llama3",
            messages=[{
                "role":"user",
                "content":f"Explain this excel data:\n\n{text}"
            }]
        )

        print("\n🤖 Excel Analysis:")
        print(r["message"]["content"])

    except Exception as e:
        print(f"❌ Excel error: {e}")


def handle_ppt_file(file_path):

    try:
        print("📽 Reading PowerPoint...")

        prs = Presentation(file_path)

        text = ""

        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + "\n"

        text = text[:4000]

        print("🧠 Sending PowerPoint content to AI...")

        r = ollama.chat(
            model="llama3",
            messages=[{
                "role":"user",
                "content":f"Summarize this presentation:\n\n{text}"
            }]
        )

        print("\n🤖 PPT Summary:")
        print(r["message"]["content"])

    except Exception as e:
        print(f"❌ PowerPoint error: {e}")

# ================= ATTACH =================
def attach_tool(file_path):
    global project_context # 🔥 ضفنا دي عشان نقدر نعدل في الذاكرة

    if not os.path.exists(file_path):
        print("❌ File not found.")
        return

    file_type = detect_file_type(file_path)

    # 🔥 السحر هنا: لو دي سكرين شوت لايف، امسح القديمة من الذاكرة الأول!
    if "live_screen" in file_path:
        project_context = [item for item in project_context if "live_screen" not in item["path"]]

    already_exists = any(item["path"] == file_path for item in project_context)

    if not already_exists:
        project_context.append({
            "path": file_path,
            "type": file_type
        })
    else:
        print("⚡ File already in context, skipped.")

    save_project_context()

# ================= LOOP =================

load_project_context()

print("\n🤖 AGENT READY (Commands: build, run, delete, job, or chat)")

while True:
    try:
        user = input().strip()
        last_user_input = user

        intent = detect_intent(user)
        if intent != "default":
            active_intent = intent
        
        if user.lower() == "exit": break
        if not user: continue

        mode = detect_mode(user)
        # 🔥 التعديل الأول: وقفنا طباعة المود هنا عشان ننضف الشاشة
        # if mode != "ATTACH":
        #     print(f"Mode: {mode}")

        if mode == "RUN":
            last = load_last_project()
            if last and os.path.exists(last):
                run_python(last)
            else:
                print("❌ No valid project history found.")

        elif mode == "DELETE":
            delete_tool(user)

        elif mode == "PROJECT":
            project_tool(user)

        elif mode == "JOB":
            job_tool(user)
        
        elif mode == "ATTACH":
            file_path = user.replace("attach", "", 1).strip()
            # print(f"📎 File attached: {file_path}")
            attach_tool(file_path)
        
        elif mode == "CLEAR":
            project_context.clear()

            if os.path.exists(CONTEXT_FILE):
                os.remove(CONTEXT_FILE)
            print("🧹 Project context cleared.")

        else:
            # 🔥 التعديل التاني: ضفنا رسالة التفكير هنا
            print("\n💭 Thinking...", flush=True)
            chat_tool(user)
            
    except KeyboardInterrupt:
        print("\nExiting...")
        break
    except Exception as e:
        print(f"❌ Critical Error: {e}")
